#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import requests
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import textwrap
import json

# === Basic Configuration ===
CSV_FILE = "2025-10-23T21-11_export.csv"
OUTPUT_FILE = "TCIA_Dataset_Slides_ollama.pptx"
OLLAMA_MODEL = "mistral"
OLLAMA_URL = "http://localhost:11434/api/generate"

# === Font Size Settings ===
TITLE_SIZE = Pt(24)
SECTION_TITLE_SIZE = Pt(18)
BODY_SIZE = Pt(13)
MIN_BODY_SIZE = Pt(10)

# === Load Data ===
df = pd.read_csv(CSV_FILE)

# === Initialize PowerPoint ===
prs = Presentation()

def wrap_text(text, width=100):
    """Wrap text to prevent overflow on slides."""
    if pd.isna(text):
        return ""
    return "\n".join(textwrap.wrap(str(text), width=width))

# === Ollama API Call ===
def ollama_generate(prompt, model="mistral"):
    """Call Ollama API and correctly parse streaming JSON responses."""
    payload = {"model": model, "prompt": prompt, "stream": True}
    response = requests.post(OLLAMA_URL, json=payload, stream=True)

    if response.status_code != 200:
        print(f"Ollama request failed: {response.status_code}")
        return ""

    full_text = ""
    for line in response.iter_lines(decode_unicode=True):
        if not line.strip():
            continue
        try:
            data = json.loads(line)
            if "response" in data:
                full_text += data["response"]
        except json.JSONDecodeError:
            continue

    return full_text.strip()

# === Main Slide Creation Function ===
def add_dataset_slide(row, index):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide_width = prs.slide_width

    # ======= Black Header Bar =======
    header_height = Inches(1.2)
    header = slide.shapes.add_shape(1, 0, 0, slide_width, header_height)
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor(0, 0, 0)

    # ======= Title =======
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(12.3), header_height - Inches(0.2))
    title_tf = title_box.text_frame
    title_tf.word_wrap = True
    title_tf.vertical_anchor = 1

    title_p = title_tf.add_paragraph()
    title_p.text = f"{index + 1}. {wrap_text(row['dataset_title'], width=53)}"
    title_p.font.size = TITLE_SIZE
    title_p.font.bold = True
    title_p.font.color.rgb = RGBColor(255, 255, 255)

    # ======= Main Content =======
    content_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12.0), Inches(6.0))
    tf = content_box.text_frame
    tf.word_wrap = True

    def add_section(title, body):
        """Add a bold section title and a text paragraph."""
        if not body or str(body).strip() == "":
            return
        p = tf.add_paragraph()
        p.text = f"{title}"
        p.font.bold = True
        p.font.size = SECTION_TITLE_SIZE
        p.space_after = Pt(4)

        p = tf.add_paragraph()
        p.text = wrap_text(str(body), width=110)
        p.font.size = BODY_SIZE
        p.space_after = Pt(10)

    # === Base Information ===
    add_section("Data:",
                f"{row['file_formats']} | Subjects: {row['n_subjects']} | Studies: {row['n_studies']} | Faces: {row['faces']}")
    add_section("Overview:", row['dataset_description'])
    add_section("Research Utility:", row['why_publish'])

    # === LLM-Generated Sections ===
    uniq_prompt = f"""
You are writing the "Uniqueness to TCIA" section for a TCIA dataset slide.
Summarize in 2–3 short sentences what makes this dataset unique compared to existing TCIA collections,
based on the following information:

Dataset title: {row['dataset_title']}
Data formats: {row['file_formats']}
Preprocessing notes: {row['preprocessing']}
Description: {row['dataset_description']}
Research goal: {row['why_publish']}
"""

    tech_prompt = f"""
You are writing the "Technical and Resource Considerations" section for a TCIA dataset slide.
Summarize in 2-3 short sentences the technical preprocessing, anonymization, and data management aspects,
based on the following information:

Preprocessing details: {row['preprocessing']}
File formats: {row['file_formats']}
Dataset title: {row['dataset_title']}
"""

    uniq_summary = ollama_generate(uniq_prompt)
    tech_summary = ollama_generate(tech_prompt)

    add_section("Uniqueness to TCIA:", uniq_summary)
    add_section("Technical and Resource Considerations:", tech_summary)

    # === Auto-adjust font size to prevent overflow ===
    total_chars = sum(len(p.text) for p in tf.paragraphs)
    if total_chars > 1800:
        for p in tf.paragraphs:
            p.font.size = MIN_BODY_SIZE

# === Generate a slide for each dataset ===
for i, row in df.iterrows():
    add_dataset_slide(row, i)

# === Save File ===
prs.save(OUTPUT_FILE)
print(f"File successfully generated: {OUTPUT_FILE}")
